"""Turn an uploaded CV into plain text. Supports PDF, DOCX, Markdown and plain text."""

import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED = (".pdf", ".docx", ".md", ".txt")


class CVError(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise CVError("This PDF is password-protected")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except PyPdfError as exc:
        raise CVError(f"Could not read this PDF: {exc}") from exc


def _from_docx(data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(BytesIO(data))
    except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as exc:
        raise CVError("Could not read this DOCX file") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _from_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise CVError("Could not decode this text file")


def extract_cv_text(filename: str, data: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in SUPPORTED:
        raise CVError(f"Unsupported file type '{extension or filename}'. Use PDF, DOCX, MD or TXT")
    if len(data) > MAX_BYTES:
        raise CVError("File is larger than 5 MB")
    reader = {".pdf": _from_pdf, ".docx": _from_docx}.get(extension, _from_text)
    text = _clean(reader(data))
    if not text:
        raise CVError(
            "No text found in this file. If it is a scanned PDF, export a text version "
            "(OCR is not supported)"
        )
    return text
