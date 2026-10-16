from io import BytesIO

import pytest
from docx import Document

from jobtrail.ingestion.cv import CVError, extract_cv_text


def docx_bytes(paragraphs, table=None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def minimal_pdf(text: str) -> bytes:
    """A tiny valid one-page PDF whose text layer says `text`."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return out


def test_markdown_and_plain_text():
    assert (
        extract_cv_text("cv.md", b"# Alex\n\n\n\n\nSkills: Python  \n")
        == "# Alex\n\nSkills: Python"
    )
    assert extract_cv_text("cv.TXT", "Café résumé".encode()) == "Café résumé"


def test_text_file_with_a_windows_encoding():
    assert extract_cv_text("cv.txt", "Café".encode("cp1252")) == "Café"


def test_docx_paragraphs_and_tables():
    data = docx_bytes(
        ["Alex Example", "Python developer"], table=[["Skill", "Level"], ["SQL", "Advanced"]]
    )
    text = extract_cv_text("cv.docx", data)
    assert "Alex Example" in text and "Python developer" in text
    assert "SQL | Advanced" in text


def test_pdf_text_layer():
    assert "Python and SQL" in extract_cv_text("cv.pdf", minimal_pdf("Python and SQL"))


@pytest.mark.parametrize(
    ("filename", "data", "message"),
    [
        ("cv.exe", b"x", "Unsupported"),
        ("noextension", b"x", "Unsupported"),
        ("cv.md", b"   \n  ", "No text found"),
        ("cv.md", b"x" * (5 * 1024 * 1024 + 1), "larger than 5 MB"),
        ("cv.pdf", b"not a pdf", "Could not read this PDF"),
        ("cv.docx", b"not a docx", "Could not read this DOCX"),
        ("cv.txt", b"\x81\x8d\x8f", "decode"),
    ],
)
def test_bad_files_get_a_clear_error(filename, data, message):
    with pytest.raises(CVError, match=message):
        extract_cv_text(filename, data)
